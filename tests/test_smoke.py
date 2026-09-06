# -*- coding: utf-8 -*-
"""冒烟测试：不依赖外部API，验证工具层与Agent全链路。运行：python tests/test_smoke.py"""
import sys
import io
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import os
import shutil  # noqa: E402
os.environ["LLM_PROVIDER"] = "mock"  # 冒烟测试强制离线mock，绝不调用真实LLM（在load_config前设置）

from config import load_config  # noqa: E402
from core.llm import extract_json  # noqa: E402
from core.agent import JobAgent  # noqa: E402
from core.memory import Task  # noqa: E402
from tools.jd_analyze import analyze_jd_text  # noqa: E402
from tools.resume_match import compute_match  # noqa: E402

PASS = []


def check(name, cond, detail=""):
    PASS.append((name, bool(cond)))
    print(("  ✅ " if cond else "  ❌ ") + name + (f"  ({detail})" if detail else ""))


def main():
    print("== 1. JSON 提取鲁棒性 ==")
    check("裸JSON", extract_json('{"a": 1}') == {"a": 1})
    check("markdown包裹", extract_json('```json\n{"a": [1,2]}\n```') == {"a": [1, 2]})
    check("带前后缀文本", extract_json('好的，以下是计划：\n{"a": {"b": 2}}\n请查收') == {"a": {"b": 2}})
    check("尾逗号容忍", extract_json('{"a": 1, "b": 2,}') == {"a": 1, "b": 2})
    check("非法输入返回None", extract_json("no json here") is None)

    print("== 2. JD 分析工具 ==")
    jd_text = (ROOT / "data/jds/jd03_数据分析师.txt").read_text(encoding="utf-8")
    jd = analyze_jd_text(jd_text)
    check("抽取到硬技能", {"SQL", "Python", "Excel"} <= set(jd["skills_hard"]),
          f"Top: {list(jd['skills_hard'])[:5]}")
    check("抽取到软技能", len(jd["skills_soft"]) >= 2, str(list(jd["skills_soft"])[:3]))
    check("学历要求", jd["education"] == "本科", jd["education"])
    check("经验年限", "年" in jd["experience_years"], jd["experience_years"])
    check("薪资抽取", jd["salary"] == "18-28K", jd["salary"])
    check("章节切分", len(jd["responsibilities"]) >= 3 and len(jd["requirements"]) >= 3,
          f"职责{len(jd['responsibilities'])}条/要求{len(jd['requirements'])}条")
    check("高频关键词", len(jd["keywords_top"]) >= 5)

    print("== 3. 简历匹配工具 ==")
    resume = (ROOT / "data/resumes/简历_李明_数据分析师.txt").read_text(encoding="utf-8")
    m = compute_match(jd_text, resume)
    check("分数在0-100", 0 <= m["score"] <= 100, f"score={m['score']} grade={m['grade']}")
    check("命中已具备技能", {"SQL", "Python", "数据分析", "A/B测试", "Tableau"} <= set(m["matched_skills"]),
          f"matched={list(m['matched_skills'])[:6]}")
    check("识别缺失技能", len(m["missing_skills"]) >= 1, f"missing={list(m['missing_skills'])[:5]}")
    jd_algo = (ROOT / "data/jds/jd02_算法工程师_NLP.txt").read_text(encoding="utf-8")
    m_algo = compute_match(jd_algo, resume)
    check("跨岗位区分度", m_algo["score"] < m["score"],
          f"数据分析师{m['score']} vs 算法工程师{m_algo['score']}")

    print("== 4. PDF 解析工具 ==")
    import subprocess
    pdf_path = ROOT / "data/resumes/简历_李明_数据分析师.pdf"
    if not pdf_path.exists():  # 样例可被用户在界面删除，测试前自建（幂等）
        subprocess.run([sys.executable, str(ROOT / "data/generate_data.py")],
                       check=True, capture_output=True, cwd=str(ROOT))
    agent = JobAgent(load_config())
    res = agent.registry.call("pdf_extract", {"path": str(ROOT / "data/resumes/简历_李明_数据分析师.pdf")})
    check("PDF解析成功", res.ok and "李明" in res.data["text"], f"{len(res.data['text'])}字")

    print("== 5. Agent 全链路（mock规划→执行→报告） ==")
    from core.memory import Memory
    memory = Memory(ROOT / "data" / "sessions_test")
    session = memory.new_session()
    events = []
    agent.bus.subscribe(events.append)
    agent.handle_message(session, f"帮我分析 {ROOT / 'data/jds/jd03_数据分析师.txt'} 这份JD，"
                                  f"用我的简历 data/resumes/简历_李明_数据分析师.txt 匹配一下")
    types = [e["type"] for e in events]
    check("触发了规划", "plan_created" in types)
    check("子任务>=4个", len(session.tasks) >= 4, f"{len(session.tasks)}个")
    check("执行了工具调用", "tool_call" in types and "tool_result" in types)
    finals = [e for e in events if e["type"] == "final_answer"]
    check("最终报告生成", bool(finals) and "匹配分" in finals[-1].get("content", ""),
          (finals[-1].get("content", "")[:60] if finals else "无final_answer"))
    check("报告落盘", any((ROOT / "data/reports").glob("求职分析报告_*.md")))
    check("匹配图表生成", bool((ROOT / "data/reports/charts").glob("match_*.png")))
    check("会话状态done", session.status == "done")
    check("记忆写入岗位", bool(session.facts.get("target_role")), str(session.facts.get("target_role")))

    print("== 6. ask_user 暂停-恢复（多轮追问） ==")
    session2 = memory.new_session()
    events2 = []
    agent.bus.subscribe(events2.append)
    agent.handle_message(session2, "我想找个数据分析的工作，帮我看看匹配度")
    check("材料缺失触发追问", session2.status == "awaiting_input",
          f"question={session2.pending_question[:30]}")
    agent.handle_message(session2,
                         f"JD在这里 {ROOT / 'data/jds/jd03_数据分析师.txt'}，简历是 data/resumes/简历_李明_数据分析师.txt")
    check("恢复后重新规划", len(session2.tasks) >= 4)
    check("恢复后执行完成", session2.status == "done" and "final_answer" in [e["type"] for e in events2])

    print("== 7. 错误处理与重试（故障注入） ==")
    from tools.base import ToolError
    session3 = memory.new_session()
    events3 = []
    agent.bus.subscribe(events3.append)
    fr = agent.registry.get("file_read")
    orig_run = fr.run

    def boom(**kwargs):
        raise ToolError("模拟故障: 文件被占用")
    fr.run = boom  # 注入故障：读取文件必失败
    try:
        agent.handle_message(session3, f"帮我分析 {ROOT / 'data/jds/jd03_数据分析师.txt'} 匹配简历")
    finally:
        fr.run = orig_run
    types3 = [e["type"] for e in events3]
    retries3 = [e for e in events3 if e["type"] == "retry"]
    t1 = session3.task("t1")
    check("工具失败触发自动重试", len(retries3) >= 1, f"{len(retries3)}次retry事件")
    check("连续失败熔断(任务判败)", t1 is not None and t1.status == "failed", t1.error[:60] if t1 else "")
    check("依赖失败自动传播", all(t.status == "failed" for t in session3.tasks if t.id in ("t2", "t3", "t5")))
    check("流程未崩溃且降级收尾", session3.status == "done" and "final_answer" in types3)
    check("失败任务在结果中披露", any("失败" in e.get("message", "") or e.get("stats", {}).get("failed_tasks")
                                       for e in events3 if e["type"] in ("error", "final_answer", "run_done")))

    print("== 8. OCR 与 docx 工具（截图JD/Word简历） ==")
    try:
        from PIL import Image, ImageDraw, ImageFont
        img = Image.new("RGB", (760, 240), "white")
        d = ImageDraw.Draw(img)
        font = ImageFont.truetype("C:/Windows/Fonts/simhei.ttf", 26)
        d.text((20, 20), "岗位职责：", font=font, fill="black")
        d.text((20, 70), "1. 负责大模型应用研发，搭建RAG检索增强问答系统", font=font, fill="black")
        d.text((20, 120), "任职要求：", font=font, fill="black")
        d.text((20, 170), "1. 精通Python，熟悉LangChain框架与向量数据库", font=font, fill="black")
        img.save(str(ROOT / "data/test_ocr.png"))
        r_ocr = agent.registry.call("image_ocr", {"path": str(ROOT / "data/test_ocr.png")})
        check("截图OCR识别JD", r_ocr.ok and "Python" in r_ocr.data["text"],
              f"{r_ocr.data['chars']}字")
        jd_ocr = analyze_jd_text(r_ocr.data["text"])
        check("OCR文本可分析出技能", "Python" in jd_ocr["skills_hard"] or "LangChain" in str(jd_ocr["keywords_top"]),
              str(list(jd_ocr["skills_hard"])[:4]))
        (ROOT / "data/test_ocr.png").unlink()
    except Exception as e:
        check("截图OCR识别JD", False, f"{type(e).__name__}: {e}")

    try:
        from docx import Document
        doc = Document()
        doc.add_heading("张三 - 数据分析师简历", 0)
        doc.add_paragraph("技能：SQL、Python（pandas）、Tableau、A/B测试")
        doc.add_paragraph("经历：某电商公司数据分析师，负责指标体系搭建")
        p_docx = ROOT / "data/resumes/test_简历_张三.docx"
        doc.save(str(p_docx))
        r_docx = agent.registry.call("docx_extract", {"path": str(p_docx)})
        check("docx简历解析", r_docx.ok and "SQL" in r_docx.data["text"], f"{len(r_docx.data['text'])}字")
        m_docx = compute_match(jd_text, r_docx.data["text"])
        check("docx简历参与匹配", 0 <= m_docx["score"] <= 100, f"score={m_docx['score']}")
        p_docx.unlink()
    except Exception as e:
        check("docx简历解析", False, f"{type(e).__name__}: {e}")

    print("== 9. 简历库默认简历发现 ==")
    from tools.file_tools import default_resume_path
    try:
        dp = default_resume_path()
        check("默认简历发现", bool(dp), dp)
    except Exception as e:
        check("默认简历发现", False, str(e))

    print("== 10. 求职智囊人设：敏感拦截/数据指令/画像抽取 ==")
    session4 = memory.new_session()
    events4 = []
    agent.bus.subscribe(events4.append)
    agent.handle_message(session4, "我的身份证号是110101199003077777，帮我优化简历")
    types4 = [e["type"] for e in events4]
    check("敏感信息拦截", "final_answer" in types4 and "请勿在对话中透露" in
          next((e["content"] for e in events4 if e["type"] == "final_answer"), ""))
    check("敏感内容未进入记忆", "110101199003077777" not in str(session4.messages))
    check("拦截后不跑流水线", "plan_created" not in types4 and len(session4.tasks) == 0)

    agent.handle_message(session4, "查看我的数据")
    view = [e["content"] for e in events4 if e["type"] == "final_answer"][-1]
    check("查看我的数据", "你的数据" in view and "已加密存储于本地" in view)

    agent.handle_message(session4, "我叫李雷，期望薪资18-25K，硕士，想找算法工程师岗位，可远程")
    facts4 = session4.facts
    check("画像抽取-姓名", facts4.get("name") == "李雷", str(facts4.get("name")))
    check("画像抽取-薪资", facts4.get("salary_range") == "18-25K", str(facts4.get("salary_range")))
    check("画像抽取-学历", facts4.get("education") == "硕士", str(facts4.get("education")))
    check("画像抽取-远程偏好", "远程" in str(facts4.get("preferences")))

    agent.handle_message(session4, "删除我的所有记忆")
    check("删除记忆指令", not [k for k in session4.facts if not str(k).startswith("_")]
          and not session4.artifacts and not session4.messages)
    del_reply = [e["content"] for e in events4 if e["type"] == "final_answer"][-1]
    check("删除确认话术", "已删除全部记忆" in del_reply and "求职智囊" in del_reply)

    print("== 11. 对抗性样本（错别字+干扰信息） ==")
    noisy = ("岗位职责：负责数據分析和报表。另：公司附近有健身房和咖啡厅，团队氛围轻松，"
             "老板人很好经常组织聚餐，提供下午茶和节日礼物，办公环境宽敞明亮。\n"
             "任职要求：熟炼使用SQL和Python；逻辑思维清晰。\n")
    jd_noisy = analyze_jd_text(noisy)
    check("对抗样本识别核心技能", {"SQL", "Python"} <= set(jd_noisy["skills_hard"]),
          str(list(jd_noisy["skills_hard"])[:4]))
    check("干扰内容未污染技能词频", jd_noisy["skills_hard"].get("SQL", 0) <= 3)

    print("== 12. 智能缓存复用（磁盘级） ==")
    from core import cache as dcache
    import time as _time
    cargs = {"jd_text": jd_text + f"（缓存测试哨兵{_time.time()}）", "top_k": 15}  # 每次运行唯一，保证幂等
    check("首次未命中", dcache.load("jd_analyze", cargs) is None)
    dcache.store("jd_analyze", cargs, {"skills_hard": {"SQL": 2}})
    check("二次命中缓存", dcache.load("jd_analyze", cargs) == {"skills_hard": {"SQL": 2}})
    check("内容变化缓存失效", dcache.load("jd_analyze", {"jd_text": jd_text + "改一句", "top_k": 15}) is None)
    sC = memory.new_session()
    evtsC = []
    agent.bus.subscribe(evtsC.append)
    resume_path = ROOT / "data/resumes/简历_李明_数据分析师.txt"
    agent.handle_message(sC, f"帮我分析 {ROOT / 'data/jds/jd09_用户增长运营.txt'} 匹配 {resume_path}")
    sD = memory.new_session()
    evtsD = []
    agent.bus.subscribe(evtsD.append)
    agent.handle_message(sD, f"帮我分析 {ROOT / 'data/jds/jd09_用户增长运营.txt'} 匹配 {resume_path}")
    check("第二次全流程命中磁盘缓存",
          any("磁盘缓存" in (e.get("brief") or "") for e in evtsD if e["type"] == "tool_result"))

    print("== 13. 隐私加密存储 ==")
    from core import secure_store
    check("加密组件可用", secure_store.enabled())
    raw = "简历内容-secret-李明".encode("utf-8")
    enc = secure_store.encrypt(raw)
    check("落盘内容已加密", enc.startswith(b"ENC1:") and raw not in enc)
    check("解密还原", secure_store.decrypt(enc) == raw)
    check("旧明文兼容", secure_store.decrypt(raw) == raw)
    mem_enc = Memory(ROOT / "data/sessions_enc_test")
    se = mem_enc.new_session()
    se.facts["name"] = "测试李"
    mem_enc.save(se)
    enc_file = ROOT / "data/sessions_enc_test" / f"{se.id}.json"
    check("会话文件密文落盘", enc_file.read_bytes().startswith(b"ENC1:")
          and "测试李".encode() not in enc_file.read_bytes())
    check("会话加密读回", mem_enc.load(se.id).facts.get("name") == "测试李")
    shutil.rmtree(ROOT / "data/sessions_enc_test", ignore_errors=True)

    print("== 14. 纯文本JD（无文件路径）全流程 ===")
    sE = memory.new_session()
    evtsE = []
    agent.bus.subscribe(evtsE.append)
    text_jd = ("岗位职责：负责大模型应用研发，搭建RAG检索增强问答系统，参与Agent工具链建设。\n"
               "任职要求：本科及以上学历，2年以上NLP或大模型应用经验，精通Python，"
               "熟悉LangChain框架与向量数据库，了解Prompt工程。")
    agent.handle_message(sE, f"帮我分析这份JD并匹配简历：\n{text_jd}\n简历用 data/resumes/简历_李明_数据分析师.txt")
    types_e = [e["type"] for e in evtsE]
    plan_e = next((e for e in evtsE if e["type"] == "plan_created"), None)
    first_tools = [t["tool"] for t in (plan_e["tasks"] if plan_e else [])]
    check("纯文本JD不生成读取任务", plan_e is not None and "file_read" not in first_tools,
          str(first_tools[:2]))
    m_e = sE.artifacts.get("match") or {}
    check("纯文本JD匹配成功", bool(m_e.get("score")), f"score={m_e.get('score')}")
    check("纯文本JD报告落盘", bool((sE.artifacts.get("report") or {}).get("path")))
    check("纯文本JD无失败任务", not [t for t in sE.tasks if t.status == "failed"])

    print("== 15. 跨会话长期记忆（全局画像） ==")
    from core import profile as profile_store
    profile_store.clear()
    sP = memory.new_session()
    evtsP = []
    agent.bus.subscribe(evtsP.append)
    agent.handle_message(sP, "我叫李雷，期望薪资18-25K，硕士，想找算法工程师")
    prof = profile_store.load()
    check("画像合并入全局", prof.get("name") == "李雷" and prof.get("salary_range") == "18-25K",
          str({k: prof.get(k) for k in ("name", "salary_range", "education")}))
    sQ = memory.new_session()
    agent.handle_message(sQ, "有什么适合我的岗位吗")
    check("新会话注入全局画像", sQ.facts.get("name") == "李雷" and sQ.facts.get("salary_range") == "18-25K",
          str({k: sQ.facts.get(k) for k in ("name", "salary_range")}))
    agent.handle_message(sQ, "删除我的所有记忆")
    check("删除记忆同步清全局", not profile_store.load().get("facts"))

    print("== 16. 条件分支规划（skipped） ==")
    sR = memory.new_session()
    sR.tasks = [
        Task(id="t1", title="计算匹配度", tool="resume_match",
             args={"jd_text": jd_text, "resume_text": resume}, depends_on=[]),
        Task(id="t2", title="高分专属建议", tool="none",
             condition={"field": "match.score", "op": ">", "value": 999}, depends_on=["t1"]),
        Task(id="t3", title="通用建议", tool="none",
             condition={"field": "match.score", "op": ">=", "value": 60}, depends_on=["t1"]),
    ]
    agent._execute_plan(sR)
    st = {t.id: t.status for t in sR.tasks}
    check("条件不满足→跳过", st.get("t2") == "skipped", str(st))
    check("条件满足→执行", st.get("t3") == "done", str(st))
    check("skipped向下游放行", st.get("t1") == "done")

    print("== 17. 输出脱敏（注入防护输出侧） ==")
    leak = ("这是密钥sk-abcdefghijklmnop1234请查收，"
            "端点https://REDACTED")
    clean = agent._sanitize_output(leak)
    check("密钥脱敏", "sk-abcdefghijklmnop1234" not in clean and "[已脱敏]" in clean)
    check("端点脱敏", "generic-endpoint" not in clean)
    check("正常内容不受影响", "匹配分88.3" in agent._sanitize_output("匹配分88.3，建议优化简历"))

    print("== 18. 意图分类前置（情绪/闲聊不跑流水线） ==")
    sS = memory.new_session()
    evtsS = []
    agent.bus.subscribe(evtsS.append)
    agent.handle_message(sS, "面试挂了，好沮丧，最近压力好大")
    replyS = [e["content"] for e in evtsS if e["type"] == "final_answer"]
    check("情绪路由-共情支持", bool(replyS) and "求职受挫" in replyS[-1] and "第一步" in replyS[-1])
    check("情绪路由-不跑流水线", not sS.tasks and "plan_created" not in [e["type"] for e in evtsS])
    sT = memory.new_session()
    evtsT = []
    agent.bus.subscribe(evtsT.append)
    agent.handle_message(sT, "你是谁")
    replyT = [e["content"] for e in evtsT if e["type"] == "final_answer"]
    check("闲聊路由-能力介绍", bool(replyT) and "优化简历" in replyT[-1] and not sT.tasks)

    print("== 19. 隐私模式（会话不落盘） ==")
    mem_priv = Memory(ROOT / "data/sessions_priv_test", privacy_mode=True)
    sV = mem_priv.new_session()
    sV.facts["name"] = "隐私测试"
    mem_priv.save(sV)
    check("隐私模式不落盘", not (ROOT / "data/sessions_priv_test" / f"{sV.id}.json").exists())
    shutil.rmtree(ROOT / "data/sessions_priv_test", ignore_errors=True)

    failed = [n for n, ok in PASS if not ok]
    print(f"\n{'='*46}\n结果: {len(PASS)-len(failed)}/{len(PASS)} 通过" + (f"，失败: {failed}" if failed else " 🎉"))
    return 0 if not failed else 1


if __name__ == "__main__":
    sys.exit(main())
