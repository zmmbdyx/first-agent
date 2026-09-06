"""文件读取类工具：PDF 解析（pdfplumber）与 txt/md 读取。"""
from pathlib import Path

from tools.base import Tool, ToolError


def _resolve(path: str) -> Path:
    if not path or not str(path).strip():
        raise ToolError("文件路径为空：请提供文件路径，或直接粘贴文本内容")
    p = Path(str(path)).expanduser()
    if not p.is_absolute():
        from config import ROOT
        p = ROOT / p
    if not p.exists():
        raise ToolError(f"文件不存在: {p}")
    if p.is_dir():
        raise ToolError(f"路径是目录而非文件: {p.name}，请提供具体文件路径")
    return p


def default_resume_path() -> str:
    """简历库默认简历：data/resumes/ 下最近修改的简历文件（txt/pdf/docx）。"""
    from config import ROOT
    resume_dir = ROOT / "data" / "resumes"
    candidates = [p for p in resume_dir.iterdir()
                  if p.suffix.lower() in (".txt", ".md", ".pdf", ".docx")]
    if not candidates:
        raise ToolError("简历库为空：请上传简历（支持 txt/md/pdf/docx）")
    return str(max(candidates, key=lambda p: p.stat().st_mtime))


class PdfExtractTool(Tool):
    name = "pdf_extract"
    description = "解析 PDF 文件（简历/JD）并提取全部文本。"
    args_desc = {"path": "PDF 文件路径"}
    cost = "低"
    avg_seconds = 0.8

    def run(self, path: str = "") -> dict:
        import io
        p = _resolve(path)
        if p.suffix.lower() != ".pdf":
            raise ToolError(f"不是 PDF 文件: {p.name}")
        try:
            import pdfplumber
            from core import secure_store
            raw = secure_store.read_bytes(p)  # 兼容加密存储
            pages = []
            with pdfplumber.open(io.BytesIO(raw)) as pdf:
                for page in pdf.pages:
                    pages.append(page.extract_text() or "")
        except ToolError:
            raise
        except Exception as e:
            raise ToolError(f"PDF 解析失败: {type(e).__name__}: {e}")
        text = "\n".join(pages).strip()
        if len(text) < 20:
            raise ToolError("PDF 中几乎无可提取文本（可能是扫描件/图片版）")
        return {"path": str(p), "pages": len(pages), "text": text[:15000]}


class FileReadTool(Tool):
    name = "file_read"
    description = "读取 txt/md 文本文件内容（如粘贴保存的 JD、简历文本）。"
    args_desc = {"path": "文件路径", "max_chars": "最大读取长度，默认12000"}
    cost = "低"
    avg_seconds = 0.1

    def run(self, path: str = "", max_chars: int = 12000) -> dict:
        from core import secure_store
        p = _resolve(path)
        if p.suffix.lower() not in (".txt", ".md", ".markdown"):
            raise ToolError(f"仅支持 txt/md，PDF请用 pdf_extract，docx请用 docx_extract，图片请用 image_ocr: {p.name}")
        text = secure_store.read_bytes(p).decode("utf-8", "ignore")
        if not text.strip():
            raise ToolError("文件内容为空")
        return {"path": str(p), "text": text[: int(max_chars or 12000)]}


class DocxExtractTool(Tool):
    name = "docx_extract"
    description = "解析 Word(docx) 简历或JD文档，提取全部文本。"
    args_desc = {"path": "docx 文件路径"}
    cost = "低"
    avg_seconds = 0.3

    def run(self, path: str = "") -> dict:
        import io
        p = _resolve(path)
        if p.suffix.lower() != ".docx":
            raise ToolError(f"不是 docx 文件: {p.name}")
        try:
            from docx import Document
            from core import secure_store
            doc = Document(io.BytesIO(secure_store.read_bytes(p)))
            parts = [para.text for para in doc.paragraphs if para.text.strip()]
            for table in doc.tables:  # 表格里的简历模块（教育/技能）也提取
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        except ToolError:
            raise
        except Exception as e:
            raise ToolError(f"docx 解析失败: {type(e).__name__}: {e}")
        text = "\n".join(parts).strip()
        if len(text) < 20:
            raise ToolError("docx 中无可提取文本")
        return {"path": str(p), "text": text[:15000]}
